import json
import threading
import unittest
from io import BytesIO
from http.client import HTTPConnection
from http.server import ThreadingHTTPServer
from urllib.parse import quote

import app
from openpyxl import load_workbook


class ParserTests(unittest.TestCase):
    def test_parse_amount_supports_russian_and_machine_formats(self):
        self.assertEqual(app.parse_amount("1 234 567,89"), 1234567.89)
        self.assertEqual(app.parse_amount("44 622 636,12"), 44622636.12)
        self.assertEqual(app.parse_amount("-37 206,75"), -37206.75)
        self.assertEqual(app.parse_amount("10000000.00"), 10000000.0)
        self.assertEqual(app.parse_amount(""), 0.0)
        self.assertEqual(app.parse_amount(None), 0.0)

    def test_parse_date_normalizes_known_input_formats(self):
        self.assertEqual(app.parse_date("20.08.2025"), "2025-08-20")
        self.assertEqual(app.parse_date("2025-03-07 00:00:00.000"), "2025-03-07")
        self.assertEqual(app.parse_date("2026-04-01"), "2026-04-01")

    def test_normalize_code_removes_separators_and_keeps_cyrillic_letters(self):
        self.assertEqual(app.normalize_code("08.3.02.97070"), "0830297070")
        self.assertEqual(app.normalize_code("13.2.01.97003"), "1320197003")
        self.assertEqual(app.normalize_code("101016105Б"), "101016105Б")

    def test_find_header_value_matches_year_independent_prefix(self):
        row = {"Лимиты ПБС 2026 год": "123"}
        self.assertEqual(app.find_header_value(row, "Лимиты ПБС"), "123")

    def test_object_group_key_uses_code_and_budget_or_normalized_name(self):
        self.assertEqual(app.object_group_key({"object_code_norm": "001", "budget": "Областной бюджет"}), "001|областной бюджет")
        self.assertEqual(app.object_group_key({"object_name": "  Объект, СКК! ", "budget": ""}), "name:объект скк|")


class DataLoadTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.store = app.STORE

    def test_loads_expected_dataset_volume_and_sources(self):
        self.assertEqual(self.store.meta["records"], 4037)
        self.assertEqual(
            set(self.store.meta["sources"]),
            {"РЧБ", "Соглашения", "ГЗ: контракты", "ГЗ: платежи", "БУАУ"},
        )

    def test_detects_expected_budgets_and_snapshots(self):
        self.assertIn("Областной бюджет Амурской области", self.store.meta["budgets"])
        self.assertIn("Бюджет г. Тынды", self.store.meta["budgets"])
        self.assertIn("2025-02-01", self.store.meta["snapshots"])
        self.assertIn("2026-04-01", self.store.meta["snapshots"])

    def test_code_filter_finds_rcb_and_agreement_records(self):
        result = app.aggregate(
            app.apply_filters(
                self.store.records,
                {"code": ["970"], "start": ["2025-01-01"], "end": ["2026-04-01"]},
            )
        )
        self.assertEqual(result["count"], 997)
        self.assertGreaterEqual(len(result["rows"]), 10)
        top_sources = result["rows"][0]["sources"]
        self.assertIn("РЧБ", top_sources)
        self.assertIn("Соглашения", top_sources)
        self.assertGreater(result["totals"]["limit"], 0)
        self.assertGreater(result["totals"]["agreement"], 0)

    def test_text_budget_and_source_filters_are_combined(self):
        filtered = app.apply_filters(
            self.store.records,
            {
                "q": ["Тынды"],
                "budget": ["Бюджет г. Тынды"],
                "source": ["БУАУ"],
                "start": ["2025-08-01"],
                "end": ["2025-12-31"],
            },
        )
        self.assertGreater(len(filtered), 0)
        self.assertTrue(all(record["source"] == "БУАУ" for record in filtered))
        self.assertTrue(all("Тынды".lower() in " ".join(str(v) for v in record.values()).lower() for record in filtered))

    def test_aggregate_has_stable_shape_for_frontend(self):
        result = app.aggregate(self.store.records)
        self.assertEqual(
            set(result.keys()),
            {"totals", "rows", "details", "timeline", "count"},
        )
        self.assertLessEqual(len(result["rows"]), 300)
        self.assertLessEqual(len(result["details"]), 500)
        self.assertTrue(result["timeline"])
        for metric in ["limit", "obligation", "cash", "agreement", "contract", "payment", "buau"]:
            self.assertIn(metric, result["totals"])

    def test_records_have_trace_fields(self):
        record = self.store.records[0]
        self.assertIn("id", record)
        self.assertIn("source_file", record)
        self.assertIn("source_row", record)
        self.assertIn("raw", record)
        self.assertTrue(record["source_file"])
        self.assertTrue(record["source_row"])

    def test_metric_selection_limits_totals(self):
        result = app.aggregate(self.store.records, ["limit", "cash"])
        self.assertEqual(set(result["totals"]), {"limit", "cash"})
        self.assertGreater(result["totals"]["limit"], 0)
        self.assertGreater(result["totals"]["cash"], 0)

    def test_templates_match_expected_code_fragments(self):
        self.assertTrue(app.matches_template({"object_code_norm": "000006105Б"}, "skk"))
        self.assertTrue(app.matches_template({"object_code_norm": "00000975"}, "kik"))
        self.assertTrue(app.matches_template({"object_code_norm": "00000978"}, "kik"))
        self.assertTrue(app.matches_template({"object_code_norm": "00000970"}, "two_thirds"))
        self.assertTrue(app.matches_template({"object_code_norm": "", "kvr": "414"}, "okv"))

    def test_kik_template_matches_current_case_data(self):
        rows = app.apply_filters(
            self.store.records,
            {"template": ["kik"], "start": ["2025-02-01"], "end": ["2026-04-02"]},
        )
        result = app.aggregate(rows)
        self.assertGreater(len(rows), 0)
        self.assertGreater(len(result["rows"]), 0)
        self.assertGreater(result["totals"]["limit"], 0)
        self.assertTrue(all(record["object_code_norm"][5:8] in {"975", "978"} for record in rows if record["object_code_norm"]))

    def test_quick_actions_include_core_scenarios(self):
        self.assertGreaterEqual(len(app.quick_actions_payload()), 6)
        for code in ("show_skk", "show_kik", "show_two_thirds", "show_okv"):
            self.assertIn(code, app.QUICK_ACTIONS)
            self.assertIn("metrics", app.QUICK_ACTIONS[code])

    def test_low_execution_post_filter_keeps_problem_rows_only(self):
        metrics = ["limit", "cash", "payment", "buau"]
        result = app.aggregate(self.store.records, metrics)
        filtered = app.apply_aggregate_post_filter(result, "low_execution", metrics)
        self.assertGreater(len(filtered["rows"]), 0)
        self.assertLessEqual(len(filtered["rows"]), len(result["rows"]))
        for row in filtered["rows"]:
            plan = float(row.get("limit") or 0) + float(row.get("obligation") or 0)
            execution = float(row.get("cash") or 0) + float(row.get("payment") or 0) + float(row.get("buau") or 0)
            self.assertGreater(plan, 0)
            self.assertTrue(execution == 0 or execution / plan < 0.25)

    def test_query_as_of_does_not_sum_monthly_rcb_snapshots(self):
        records = [
            {"source": "РЧБ", "record_kind": "rcb_snapshot", "snapshot": "2025-01-01", "object_code_norm": "1", "object_code": "1", "object_name": "obj", "budget": "", "limit": 10, "obligation": 0, "cash": 0, "agreement": 0, "contract": 0, "payment": 0, "buau": 0},
            {"source": "РЧБ", "record_kind": "rcb_snapshot", "snapshot": "2025-02-01", "object_code_norm": "1", "object_code": "1", "object_name": "obj", "budget": "", "limit": 20, "obligation": 0, "cash": 0, "agreement": 0, "contract": 0, "payment": 0, "buau": 0},
        ]
        result = app.aggregate(app.select_as_of(records, "2025-02-15", {}))
        self.assertEqual(result["totals"]["limit"], 20)

    def test_rcb_2026_limit_and_obligation_columns_are_loaded(self):
        result = app.query_as_of({"date": ["2026-04-01"], "template": ["all"]})
        self.assertGreater(result["totals"]["limit"], 0)
        self.assertGreater(result["totals"]["obligation"], 0)

    def test_agreements_are_not_duplicated_across_monthly_snapshots(self):
        records = [
            {"source": "Соглашения", "record_kind": "agreement_snapshot", "snapshot": "2025-01-01", "document_id": "a1", "object_code_norm": "1", "object_code": "1", "object_name": "obj", "budget": "", "limit": 0, "obligation": 0, "cash": 0, "agreement": 50, "contract": 0, "payment": 0, "buau": 0},
            {"source": "Соглашения", "record_kind": "agreement_snapshot", "snapshot": "2025-02-01", "document_id": "a1", "object_code_norm": "1", "object_code": "1", "object_name": "obj", "budget": "", "limit": 0, "obligation": 0, "cash": 0, "agreement": 60, "contract": 0, "payment": 0, "buau": 0},
        ]
        result = app.aggregate(app.select_as_of(records, "2025-02-15", {}))
        self.assertEqual(result["totals"]["agreement"], 60)

    def test_reporting_dates_exclude_payment_only_dates(self):
        payload = app.reporting_dates_payload()
        dates = [item["date"] for item in payload]
        self.assertIn("2026-04-01", dates)
        self.assertNotIn("2026-04-02", dates)

    def test_compare_uses_as_of_semantics(self):
        result = app.compare_periods({"base": ["2025-02-01"], "target": ["2026-04-01"], "template": ["skk"], "metrics": ["limit,cash"]})
        self.assertEqual(result["view"], "as_of")
        self.assertEqual(result["available_dates"], app.STORE.meta["reporting_dates"])
        self.assertTrue(result["rows"])

    def test_problem_filters_no_cash_no_payments_no_documents(self):
        records = [
            {"source": "РЧБ", "record_kind": "rcb_snapshot", "snapshot": "2026-04-01", "object_code_norm": "1", "object_code": "1", "object_name": "no cash", "budget": "", "limit": 100, "obligation": 0, "cash": 0, "agreement": 10, "contract": 0, "payment": 1, "buau": 0},
            {"source": "РЧБ", "record_kind": "rcb_snapshot", "snapshot": "2026-04-01", "object_code_norm": "2", "object_code": "2", "object_name": "no docs", "budget": "", "limit": 100, "obligation": 0, "cash": 1, "agreement": 0, "contract": 0, "payment": 0, "buau": 0},
            {"source": "Соглашения", "record_kind": "agreement_snapshot", "snapshot": "2026-04-01", "object_code_norm": "3", "object_code": "3", "object_name": "no pay", "budget": "", "limit": 0, "obligation": 0, "cash": 0, "agreement": 10, "contract": 0, "payment": 0, "buau": 0},
        ]
        result = app.aggregate(records)
        self.assertEqual(len(app.apply_aggregate_post_filter(result, "no_cash", app.METRIC_KEYS)["rows"]), 1)
        self.assertEqual(len(app.apply_aggregate_post_filter(result, "no_documents", app.METRIC_KEYS)["rows"]), 1)
        self.assertEqual(len(app.apply_aggregate_post_filter(result, "no_payments", app.METRIC_KEYS)["rows"]), 1)

    def test_pipeline_fields_are_returned_for_rows(self):
        result = app.query_as_of({"date": ["2026-04-01"], "template": ["skk"]})
        row = result["rows"][0]
        self.assertIn("pipeline", row)
        self.assertIn("plan", row["pipeline"])
        self.assertIn("problem_reasons", row)
        self.assertIn("object_key", row)
        self.assertIn("match_confidence", row)

    def test_risk_score_and_level_are_added_to_rows(self):
        result = app.query_as_of({"date": ["2026-04-01"], "template": ["skk"]})
        row = result["rows"][0]
        self.assertIn("risk_score", row)
        self.assertIn("risk_level", row)
        self.assertIn("risk_label", row)
        self.assertIn("risk_explanation", row)

    def test_attention_summary_returns_human_bullets(self):
        result = app.query_as_of({"date": ["2026-04-01"], "template": ["skk"]})
        summary = result["attention_summary"]
        self.assertEqual(summary["title"], "Что требует внимания")
        self.assertTrue(summary["bullets"])
        text = " ".join(summary["bullets"])
        self.assertNotIn("problem_reasons", text)
        self.assertNotIn("pipeline", text)

    def test_problem_rows_are_sorted_by_risk(self):
        result = app.query_as_of({
            "date": ["2026-04-01"],
            "template": ["skk"],
            "post_filter": ["execution_problems"],
        })
        scores = [row["risk_score"] for row in result["rows"]]
        self.assertEqual(scores, sorted(scores, reverse=True))

    def test_compare_insights_are_returned(self):
        result = app.compare_periods({
            "base": ["2025-02-01"],
            "target": ["2026-04-01"],
            "template": ["skk"],
        })
        self.assertIn("compare_insights", result)
        self.assertIn("bullets", result["compare_insights"])

    def test_object_detail_returns_first_query_row(self):
        result = app.query_as_of({"date": ["2026-04-01"], "template": ["skk"]})
        row = result["rows"][0]
        payload = app.object_detail({"date": ["2026-04-01"], "template": ["skk"], "object_key": [row["object_key"]]})
        self.assertEqual(payload["object_key"], row["object_key"])
        self.assertIn("pipeline", payload)
        self.assertIn("documents", payload)

    def test_object_detail_includes_risk(self):
        query = app.query_as_of({"date": ["2026-04-01"], "template": ["skk"]})
        key = query["rows"][0]["object_key"]
        detail = app.object_detail({"date": ["2026-04-01"], "template": ["skk"], "object_key": [key]})
        self.assertIn("risk_score", detail)
        self.assertIn("risk_label", detail)

    def test_readiness_summary_returns_checks(self):
        payload = app.readiness_response({"date": ["2026-04-01"], "template": ["skk"]})
        self.assertEqual(payload["date"], "2026-04-01")
        self.assertTrue(payload["checks"])
        self.assertIn("summary", payload)

    def test_as_of_timeline_uses_reporting_dates_and_matches_totals(self):
        result = app.query_as_of({"date": ["2026-04-01"], "template": ["skk"], "metrics": ["limit,cash"]})
        dates = [point["date"] for point in result["timeline"]]
        self.assertTrue(dates)
        self.assertTrue(all(date in app.STORE.meta["reporting_dates"] for date in dates))
        self.assertEqual(dates[-1], "2026-04-01")
        self.assertEqual(result["timeline"][-1]["limit"], result["totals"]["limit"])
        self.assertEqual(result["timeline"][-1]["cash"], result["totals"]["cash"])

    def test_empty_readiness_marks_empty_result(self):
        payload = app.readiness_response({"date": ["2026-04-01"], "template": ["skk"], "q": ["zzzz-no-data"]})
        empty = next(check for check in payload["checks"] if check["code"] == "empty_result")
        self.assertIn(empty["status"], {"warn", "bad"})

    def test_quick_actions_use_as_of_date(self):
        self.assertEqual(app.QUICK_ACTIONS["execution_problems"]["post_filter"], "execution_problems")
        self.assertEqual(app.default_date_range()[1], app.STORE.meta["reporting_dates"][-1])

    def test_assistant_rule_based_core_intents(self):
        skk = app.assistant_rule_based("Покажи СКК", {})
        self.assertEqual(skk["action"]["template"], "skk")
        self.assertEqual(skk["alternatives"][0]["label"], "Показать все данные")
        self.assertEqual(skk["alternatives"][0]["action"]["q"], "")

        city = app.assistant_rule_based("6105 Благовещенск", {})
        self.assertEqual(city["action"]["template"], "skk")
        self.assertIn("Благовещенск", city["action"]["q"])
        self.assertEqual(city["alternatives"][0]["label"], "Искать во всех данных")
        self.assertTrue(city["alternatives"][0]["action"]["reset_scope"])
        self.assertEqual(city["alternatives"][0]["action"]["code"], "")
        self.assertEqual(city["alternatives"][0]["action"]["budget"], "")
        self.assertEqual(city["alternatives"][0]["action"]["source"], "")

        compare = app.assistant_rule_based("сравни СКК", {})
        self.assertEqual(compare["intent"], "run_compare")
        self.assertEqual(compare["action"]["mode"], "compare")

        explain = app.assistant_rule_based("что такое БО", {})
        self.assertEqual(explain["intent"], "explain_metric")

    def test_rag_loader_reads_markdown_documents(self):
        documents = app.load_rag_documents()
        self.assertGreaterEqual(len(documents), 6)
        self.assertTrue(any("БО" in item["content"] for item in documents))


class HttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), app.Handler)
        cls.port = cls.server.server_address[1]
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=5)

    def request(self, path, method="GET", payload=None):
        conn = HTTPConnection("127.0.0.1", self.port, timeout=10)
        try:
            body = None
            headers = {}
            if payload is not None:
                body = json.dumps(payload).encode("utf-8")
                headers["Content-Type"] = "application/json"
            conn.request(method, path, body=body, headers=headers)
            response = conn.getresponse()
            body = response.read()
            return response.status, response.getheader("Content-Type"), body
        finally:
            conn.close()

    def test_meta_endpoint_returns_json(self):
        status, content_type, body = self.request("/api/meta")
        self.assertEqual(status, 200)
        self.assertIn("application/json", content_type)
        payload = json.loads(body.decode("utf-8"))
        self.assertEqual(payload["records"], 4037)
        self.assertIn("РЧБ", payload["sources"])

    def test_query_endpoint_returns_filtered_result(self):
        status, content_type, body = self.request("/api/query?code=6105&source=%D0%A0%D0%A7%D0%91")
        self.assertEqual(status, 200)
        self.assertIn("application/json", content_type)
        payload = json.loads(body.decode("utf-8"))
        self.assertGreater(payload["count"], 0)
        self.assertTrue(payload["rows"])
        self.assertTrue(all("РЧБ" in row["sources"] for row in payload["rows"]))

    def test_static_index_is_served(self):
        status, content_type, body = self.request("/")
        self.assertEqual(status, 200)
        self.assertIn("text/html", content_type)
        self.assertIn("Простая аналитика расходов", body.decode("utf-8"))


    def test_catalog_quality_trace_and_compare_endpoints_return_json(self):
        for path in (
            "/api/catalog/templates",
            "/api/catalog/metrics",
            "/api/catalog/quick-actions",
            "/api/catalog/dates",
            "/api/catalog/sources",
            "/api/catalog/budgets",
            "/api/catalog/objects?q=6105",
            "/api/quality",
            "/api/readiness?view=as_of&date=2026-04-01&template=skk",
            "/api/compare?base=2025-02-01&target=2026-04-01&template=skk&metrics=limit,cash",
        ):
            status, content_type, body = self.request(path)
            self.assertEqual(status, 200, path)
            self.assertIn("application/json", content_type)
            self.assertIsNotNone(json.loads(body.decode("utf-8")))

        record_id = app.STORE.records[0]["id"]
        status, content_type, body = self.request(f"/api/trace?id={record_id}")
        self.assertEqual(status, 200)
        payload = json.loads(body.decode("utf-8"))
        self.assertEqual(payload["id"], record_id)
        self.assertIn("raw", payload)
        self.assertIn("human_summary", payload)

    def test_object_and_excel_export_endpoints(self):
        query = app.query_as_of({"date": ["2026-04-01"], "template": ["skk"]})
        object_key = query["rows"][0]["object_key"]
        status, content_type, body = self.request(f"/api/object?date=2026-04-01&template=skk&object_key={quote(object_key)}")
        self.assertEqual(status, 200)
        self.assertIn("application/json", content_type)
        payload = json.loads(body.decode("utf-8"))
        self.assertEqual(payload["object_key"], object_key)

        status, content_type, body = self.request("/api/export.xlsx?date=2026-04-01&template=skk")
        self.assertEqual(status, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", content_type)
        workbook = load_workbook(BytesIO(body), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Выводы", "Итоги", "Объекты", "Проблемы", "Исходные строки", "Методика"])
        self.assertIn("Выводы", workbook.sheetnames)
        method_text = "\n".join(str(row[0] or "") for row in workbook["Методика"].iter_rows(values_only=True))
        self.assertIn("последний месячный срез", method_text)

    def test_assistant_endpoint_returns_rule_based_json_without_groq(self):
        status, content_type, body = self.request(
            "/api/assistant",
            method="POST",
            payload={"message": "Покажи СКК", "context": {"mode": "slice", "template": "all"}},
        )
        self.assertEqual(status, 200)
        self.assertIn("application/json", content_type)
        payload = json.loads(body.decode("utf-8"))
        self.assertIn(payload["mode"], {"rule_based", "llm"})
        self.assertEqual(payload["action"]["template"], "skk")

    def test_quick_actions_run_against_public_api(self):
        status, _, body = self.request("/api/catalog/quick-actions")
        self.assertEqual(status, 200)
        actions = json.loads(body.decode("utf-8"))
        self.assertGreaterEqual(len(actions), 6)
        for action in actions:
            metrics = ",".join(action.get("metrics", []))
            if action["mode"] == "compare":
                path = f"/api/compare?template={action['template']}&base=2025-02-01&target=2026-04-01&metrics={metrics}"
                status, content_type, body = self.request(path)
                self.assertEqual(status, 200, action["code"])
                payload = json.loads(body.decode("utf-8"))
                self.assertIn("rows", payload)
            else:
                post_filter = f"&post_filter={action['post_filter']}" if action.get("post_filter") else ""
                path = f"/api/query?template={action['template']}&start=2025-02-01&end=2026-04-01&metrics={metrics}{post_filter}"
                status, content_type, body = self.request(path)
                self.assertEqual(status, 200, action["code"])
                payload = json.loads(body.decode("utf-8"))
                self.assertIn("rows", payload)
                self.assertIn("totals", payload)


class StaticFilesTests(unittest.TestCase):
    def test_frontend_assets_exist_and_reference_api(self):
        index = (app.STATIC_DIR / "index.html").read_text(encoding="utf-8")
        script = (app.STATIC_DIR / "app.js").read_text(encoding="utf-8")
        styles = (app.STATIC_DIR / "styles.css").read_text(encoding="utf-8")

        self.assertIn("/static/app.js", index)
        self.assertIn("/api/meta", script)
        self.assertIn("/api/query", script)
        self.assertIn("/api/compare", script)
        self.assertIn("/api/trace", script)
        self.assertIn("/api/readiness", script)
        self.assertIn("/api/object", script)
        self.assertIn("/api/export.xlsx", script)
        self.assertIn("/api/catalog/quick-actions", script)
        self.assertIn("/api/assistant", script)
        self.assertIn("grid-template-columns", styles)

    def test_vue_frontend_is_declarative(self):
        index = (app.STATIC_DIR / "index.html").read_text(encoding="utf-8")
        script = (app.STATIC_DIR / "app.js").read_text(encoding="utf-8")

        self.assertIn('id="app"', index)
        self.assertIn("vue.global.prod.js", index)
        self.assertIn("/static/app.js", index)
        for marker in (
            "createApp",
            'mount("#app")',
            "loadInitialData",
            "loadData",
            "loadCompare",
            "openTrace",
            "exportCsv",
            "exportExcel",
            "openObject",
            "quickActions",
            "smartInput",
            "assistant",
            "applyQuickAction",
            "buildSmartSuggestions",
            "askAssistant",
            "applyAssistantAction",
            "resultNarrative",
            "simpleRows",
            "problemRows",
        ):
            self.assertIn(marker, script)
        for marker in (
            "Спросить помощника",
            "Короткий вывод",
            "Объекты",
            "Проблемы",
        ):
            self.assertIn(marker, index)
        for legacy_marker in (
            "document.querySelector",
            "innerHTML =",
            "addEventListener",
        ):
            self.assertNotIn(legacy_marker, script)


if __name__ == "__main__":
    unittest.main()
